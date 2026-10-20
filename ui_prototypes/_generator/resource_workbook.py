# -*- coding: utf-8 -*-
"""Build-time formula workbook for a resource-engine company (paper section 6 gap 9, 20 Oct 2026).

The companion to ``engine_workbook.py`` for companies whose ``normalised_baseline.engine``
is ``resource`` (D-96). A mine-by-mine depletion run-off (D-97) on the Assumptions-sheet
discipline (standing rule 1): every yellow cell on ``Assumptions`` is a DATA input read
from the data files; every other cell on every other sheet is an Excel formula that links
back to it. Nothing is computed in Python and pasted. The engine is not run here at all;
the tie is asserted by ``tests/dcf/test_evn_workbook_tie.py`` against an oracle that
LibreOffice recalculates (``tests/dcf/golden/_recalc_generated_workbooks.py``).

Sheets
    Assumptions   the inputs: valuation date and stub, WACC build, equity bridge, inventory
                  basis selector, guidance per mine (low / high typed; mid derived), project
                  uplifts (derived mines), the D-104 cost-price link as the archetype
                  baseline and the company override on separate rows (methodology section
                  11), the working-capital derivation, and each scenario's series anchors.
    Prices        per scenario: CPI and cost indices, AUD/USD, real USD and nominal A$
                  prices by period, interpolated between the anchors (the D-37 convention),
                  and the D-104 multiplier.
    Run-off <X>   one sheet per scenario: each mine's remaining inventory, production,
                  revenue, costs and capital by period, the group build to FCFF, mid-period
                  discounting, PV explicit and PV run-off, the equity bridge, value per
                  share, run-off share and life.
    Scenarios     the six values side by side, against the market price.

Names no company, mine or scenario: the mines are whatever the guidance lists, the
series are whatever the archetype names. ``build_resource_workbook_bytes(cid, arch)``.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(_ROOT / "src"))
sys.path.insert(0, str(_ROOT / "ui_prototypes" / "_generator"))

from engine_workbook import (  # noqa: E402
    BLUEFONT, BOLD, HDR, NOTE, NUM0, NUM1, PCT2, SUB, YELLOW, SCEN, _working_capital_block,
)
from vcc_valuations.translator import load_inputs, scenario_equilibrium_year  # noqa: E402
from vcc_valuations.resource_translator import (  # noqa: E402
    INVENTORY_BASES, _framework, build_resource_inputs_from_data, cost_price_link_params,
)

SIDS = [sid for sid, _ in SCEN]
NAMES = dict(SCEN)
SHORT = {"orderly_convergence": "OC", "muddle_through": "MT", "ai_productivity_lag": "AI",
         "fragmentation": "Frag", "disorderly_climate_crystallisation": "DCC", "stagflation_persists": "Stag"}
NUM2 = "#,##0.00"
NUM4 = "0.0000"
# Inventory basis selector: the order the engine's INVENTORY_BASES declares, numbered from 1.
BASIS_ORDER = list(INVENTORY_BASES)  # stated_life, reserves, life_extension
_KOZ_PER_MOZ = 1.0e3   # ssot-allow: unit conversion
_THOUSANDS = 1.0e3     # ssot-allow: koz x A$/oz -> A$ m
_EXHAUSTED = 1.0e-9    # ssot-allow: the engine's numerical tolerance, mirrored
_BIG = 1.0e9           # ssot-allow: sentinel for "no metal at this mine" in a MIN()
_SPARE_PERIODS = 2     # ssot-allow: the translator's own allowance past the longest life


def _col(i):
    return get_column_letter(i)


def _mid(b):
    return (b["low"] + b["high"]) / 2.0


def _periods_needed(inputs: dict) -> int:
    """Columns the run-off needs on the LONGEST inventory basis, so the basis selector can
    be flexed in Excel without running off the sheet. A layout width, not a value."""
    longest = 0.0
    for basis in INVENTORY_BASES:
        ri = build_resource_inputs_from_data(inputs, "muddle_through", basis=basis)
        for m in ri.mines:
            lives = [m.inventory_gold_koz / m.gold_koz_pa if m.gold_koz_pa else 0.0,
                     m.inventory_copper_kt / m.copper_kt_pa if m.copper_kt_pa else 0.0]
            longest = max(longest, m.start_period + max(lives))
    return int(longest) + _SPARE_PERIODS


def _series_anchors(inputs: dict, variable: str):
    for mv in inputs["scenario"].macro_variables:
        if mv.variable == variable and mv.time_series:
            return sorted((p.year, p.value) for p in mv.time_series)
    raise KeyError(f"{inputs['scenario'].id}: no time series for {variable!r}")


def _escalation_mid(inputs: dict, sid: str) -> float:
    matrix = inputs.get("matrix")
    for entry in (matrix.matrix if matrix else []):
        if entry.scenario == sid and "real_cost_escalation" in entry.drivers:
            q = entry.drivers["real_cost_escalation"].quantified_override
            return q.mid if q else 0.0
    return 0.0


def build_resource_workbook_bytes(cid: str, arch: str) -> bytes:
    by_sid = {sid: load_inputs(_ROOT, sid, arch, cid) for sid in SIDS}
    inp = by_sid["muddle_through"]
    fin = inp["financials"]
    craw = inp["company_raw"]
    cp = craw["company_position"]
    nb = craw["normalised_baseline"]
    if nb.get("engine") != "resource":
        raise ValueError(f"{cid}: normalised_baseline.engine is not 'resource'")
    short = cp["ticker"]
    ccy = cp["functional_currency"]
    g = fin["guidance_fy27"]
    mines = list(g["mines"])
    uplifts = list(nb.get("project_uplifts") or [])
    fw = _framework(inp)
    link_arch = fw.cost_price_link.model_dump() if fw.cost_price_link is not None else None
    link_co = nb.get("cost_price_link") or {}
    series = {"gold": next(c.price_series for c in fw.commodities if c.role.value == "primary"),
              "copper": next(c.price_series for c in fw.commodities if c.role.value == "by_product"),
              "fx": fw.fx_series, "cpi": fw.inflation_series}
    stream = fin.get("northparkes_stream")
    runoff = list(nb["project_capex_runoff"])
    NP = _periods_needed(inp)
    ncols_mines = len(mines) + len(uplifts)

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    R = {}  # key -> absolute reference on Assumptions

    # ======================= Assumptions =======================
    ws = wb.active
    ws.title = "Assumptions"
    ws.sheet_properties.tabColor = "FFD966"
    r = [1]

    def head(t):
        ws.cell(r[0], 1, t).font = HDR; r[0] += 1

    def sub(t):
        ws.cell(r[0], 1, t).font = SUB; r[0] += 1

    def note(t):
        ws.cell(r[0], 1, t).font = NOTE; r[0] += 1

    def line(label, value, key, fmt=None, col=2):
        """One yellow input in column ``col`` (B by default) with its label in A."""
        ws.cell(r[0], 1, label)
        c = ws.cell(r[0], col, value); c.fill = YELLOW; c.font = BLUEFONT
        if fmt:
            c.number_format = fmt
        R[key] = f"'Assumptions'!${_col(col)}${r[0]}"
        r[0] += 1

    def derived(label, formula, key, fmt=None, bold=False):
        c0 = ws.cell(r[0], 1, label)
        c = ws.cell(r[0], 2, formula)
        if fmt:
            c.number_format = fmt
        if bold:
            c0.font = BOLD; c.font = BOLD
        R[key] = f"'Assumptions'!$B${r[0]}"
        r[0] += 1

    head(f"{short} — Assumptions (yellow cells are the only inputs; mine-by-mine depletion run-off, D-96/D-97)")
    note("Every other sheet links here by formula. Sourced from the data files the production resource engine reads; nothing is typed from a result.")
    r[0] += 1

    sub("Valuation date and horizon")
    line("Valuation date", nb["valuation_date"], "vdate", "d mmm yyyy")
    line("Stub, years to the June year-end", nb["stub_years"], "f0", NUM4)
    line("Explicit horizon, years (D-71)", int(nb["horizon_years"]), "H", NUM0)
    derived("First fiscal year (June year-end: year + 1 when the date is after June)",
            f"=YEAR({R['vdate']})+IF(MONTH({R['vdate']})>6,1,0)", "fy1", "0")
    r[0] += 1

    sub("Market reference and shares")
    line(f"Market price ({ccy})", fin["market_data"]["price_aud"], "mkt", NUM2)
    line("Market reference date", fin["market_data"]["price_date"], "mkt_date", "d mmm yyyy")
    line("Shares outstanding (m), at the net-debt anchor date", fin["share_statistics"]["shares_outstanding"], "shares", NUM1 + "00")
    r[0] += 1

    sub("WACC build (D-01 single rate; D-101 held)")
    obs = fin["wacc_observed_inputs"]
    wbld = nb["wacc_build"]
    line("Risk-free rate Rf (10-year bond)", obs["risk_free_rate"], "rf", PCT2)
    line("Risk-free rate date", obs["risk_free_rate_date"], "rf_date", "d mmm yyyy")
    line("Equity risk premium ERP", wbld["equity_risk_premium"], "erp", PCT2)
    line("Beta (peer judgement; observed peer betas awaited)", wbld["beta"], "beta", NUM2)
    line("Cost of debt spread over Rf", wbld["cost_of_debt_spread"], "kd_spread", PCT2)
    line("Tax rate", wbld["tax_rate"], "tax_wacc", PCT2)
    line(f"Debt at market value ({ccy} m; notes at their swapped value)", fin["debt"]["uspp_hedged_aud_m"], "debt_mv", NUM0)
    derived("→ Cost of equity Rf + beta x ERP", f"={R['rf']}+{R['beta']}*{R['erp']}", "ke", PCT2)
    derived("→ Pre-tax cost of debt Rf + spread", f"={R['rf']}+{R['kd_spread']}", "kd", PCT2)
    derived("→ After-tax cost of debt", f"={R['kd']}*(1-{R['tax_wacc']})", "kd_at", PCT2)
    derived(f"→ Equity at market value ({ccy} m) = price x shares", f"={R['mkt']}*{R['shares']}", "eq_mv", NUM0)
    derived("→ E / (D + E), spot weights", f"={R['eq_mv']}/({R['eq_mv']}+{R['debt_mv']})", "we", PCT2)
    derived("→ WACC", f"={R['we']}*{R['ke']}+(1-{R['we']})*{R['kd_at']}", "wacc", PCT2, bold=True)
    r[0] += 1

    sub("Equity bridge (D-100: anchored at the latest quarterly net cash, both to one date)")  # ssot-allow: decision id
    dm = fin["derived_metrics"]
    bs = fin["latest_balance_sheet"]
    eb = nb["equity_bridge"]
    line(f"Net cash at the anchor ({ccy} m)", dm["net_cash_at_anchor"], "net_cash", NUM0)
    line(f"Final dividend declared and unpaid at the anchor ({ccy} m)", dm["final_dividend_unpaid_at_anchor"], "div_unpaid", NUM0)
    line("Balance-sheet date for the claims below", bs["date"], "bs_date", "d mmm yyyy")
    line(f"Rehabilitation provision ({ccy} m)", bs["rehabilitation_provision"], "rehab", NUM1)
    line("Deduct rehabilitation provision? (1 = yes, 0 = no)", int(bool(eb["deduct_rehabilitation_provision"])), "f_rehab", "0")
    line(f"Stream deferred revenue ({ccy} m; modelled in the mine cash flows)", bs["stream_deferred_revenue"], "stream_dr", NUM1)
    line("Deduct stream deferred revenue? (1 = yes, 0 = no)", int(bool(eb["deduct_stream_deferred_revenue"])), "f_stream", "0")
    line(f"Lease liabilities ({ccy} m)", bs["lease_liabilities"], "leases_in", NUM1)
    line("Deduct leases as debt? (1 = yes, 0 = no; Approach A)", int(bool(eb["deduct_lease_liabilities"])), "f_lease", "0")
    derived(f"→ Net debt at valuation ({ccy} m) = -net cash + unpaid dividend", f"=-{R['net_cash']}+{R['div_unpaid']}", "net_debt", NUM0)
    derived(f"→ Rehabilitation and other claims ({ccy} m)", f"={R['rehab']}*{R['f_rehab']}+{R['stream_dr']}*{R['f_stream']}", "claims", NUM1)
    derived(f"→ Leases deducted ({ccy} m)", f"={R['leases_in']}*{R['f_lease']}", "leases", NUM1)
    r[0] += 1

    sub("Group operating inputs (FY27 guidance midpoints; low / high typed, mid derived)")
    pb = g["price_basis"]
    line(f"Guidance price basis — gold ({ccy}/oz)", pb["gold_aud_oz"], "pb_g", NUM0)
    line(f"Guidance price basis — copper ({ccy}/t)", pb["copper_aud_t"], "pb_c", NUM0)
    for key, label, blk in (("corp", f"Corporate cost ({ccy}/oz of gold)", g["corporate_aisc_aud_oz"]),
                            ("green", f"Greenfields exploration ({ccy} m pa)", g["greenfields_exploration"])):
        line(label + " — low", blk["low"], key + "_lo", NUM1)
        line(label + " — high", blk["high"], key + "_hi", NUM1)
        derived("→ " + label + " — mid", f"=({R[key + '_lo']}+{R[key + '_hi']})/2", key, NUM1)
    line("Statutory tax rate", fin["statutory_tax_rate"], "tax", PCT2)
    ry = fin["royalties_fy26"]
    line(f"Royalty, FY26 ({ccy} per sold ounce)", ry["royalty_aud_per_sold_oz"], "ry_oz", NUM0)
    line("Gold sold, FY26 (oz)", ry["gold_sold_oz"], "ry_sold", NUM0)
    line(f"Achieved gold price, FY26 ({ccy}/oz)", ry["achieved_gold_aud_oz"], "ry_pg", NUM0)
    line("Copper produced, FY26 (t)", ry["copper_produced_t"], "ry_cu", NUM0)
    line(f"Achieved copper price, FY26 ({ccy}/t)", ry["achieved_copper_aud_t"], "ry_pc", NUM0)
    derived("→ Royalty rate, one group rate on metal revenue (derived, paper section 6 gap 5)",
            f"={R['ry_oz']}*{R['ry_sold']}/({R['ry_sold']}*{R['ry_pg']}+{R['ry_cu']}*{R['ry_pc']})", "roy", "0.000%")
    r[0] += 1

    _working_capital_block(ws, r, R, nb["working_capital_intensity"], fin["working_capital_history"])
    R["wci"] = R.pop("wc_int")
    r[0] += 1

    sub("Inventory basis (D-98): 1 = stated life of mine, 2 = Ore Reserves only, 3 = stated life + extension")
    line("Inventory basis selected (1, 2 or 3)", BASIS_ORDER.index(nb["inventory"]["basis"]) + 1, "basis", "0")
    line("Life extension, years per mine (basis 3 only)", int(nb["inventory"].get("life_extension_years", 0)), "ext", "0")
    line("Metallurgical recovery applied to Ore Reserves (basis 2)", nb["inventory"]["default_metallurgical_recovery"], "rec", PCT2)
    res = fin["ore_reserves_2025_12_31"]
    line(f"Ore Reserve price basis — gold ({ccy}/oz; disclosure only)", res["reserve_price_aud"]["gold_oz"], "res_pg", NUM0)
    line(f"Ore Reserve price basis — copper ({ccy}/t; disclosure only)", res["reserve_price_aud"]["copper_t"], "res_pc", NUM0)
    r[0] += 1

    sub("Project capital run-off (D-98): share of the FY27 major project capital in each period, stub first")
    ws.cell(r[0], 1, "Run-off fractions by period (period 0 = stub)")
    for j, x in enumerate(runoff):
        c = ws.cell(r[0], 2 + j, x); c.fill = YELLOW; c.font = BLUEFONT; c.number_format = PCT2
        R[f"runoff{j}"] = f"'Assumptions'!${_col(2 + j)}${r[0]}"
    r[0] += 1
    r[0] += 1

    sub("Cost-price link (D-104): industry-archetype baseline and company override on separate rows; the applied row is derived")
    ws.cell(r[0], 2, "Archetype").font = BOLD
    ws.cell(r[0], 3, "Company override (blank = none)").font = BOLD
    ws.cell(r[0], 4, "Applied").font = BOLD
    r[0] += 1
    link_rows = [("elasticity_per_10pct", "Elasticity, % cost move per 10% real price move", NUM1),
                 ("lag_years", "Lag, years", "0"),
                 ("cap_above_base_pct", "Cap above the base cost path, %", NUM1),
                 ("symmetric", "Symmetric? (1 = costs ease when the price falls, 0 = upward only)", "0")]
    for key, label, fmt in link_rows:
        ws.cell(r[0], 1, label)
        a = link_arch[key] if link_arch is not None else 0
        if isinstance(a, bool):
            a = int(a)
        c = ws.cell(r[0], 2, a); c.fill = YELLOW; c.font = BLUEFONT; c.number_format = fmt
        o = link_co.get(key)
        if isinstance(o, bool):
            o = int(o)
        c = ws.cell(r[0], 3, o); c.fill = YELLOW; c.font = BLUEFONT; c.number_format = fmt
        ws.cell(r[0], 4, f'=IF(ISNUMBER(C{r[0]}),C{r[0]},B{r[0]})').number_format = fmt
        R[f"link_{key}"] = f"'Assumptions'!$D${r[0]}"
        r[0] += 1
    link_on = 1 if link_arch is not None else 0
    line("Link declared by the archetype? (1 = yes, 0 = no link)", link_on, "link_on", "0")
    r[0] += 1

    # ---- per-mine guidance table: mines as columns ----
    sub("Mines (FY27 guidance; low / high typed, mid derived) and project uplifts as derived mines (method paper section 5.3)")
    hdr = r[0]
    ws.cell(hdr, 1, "Item").font = BOLD
    for j, name in enumerate(mines):
        c = ws.cell(hdr, 3 + j, name.replace("_", " ").title()); c.font = BOLD; c.alignment = Alignment(horizontal="right", wrap_text=True)
    for j, u in enumerate(uplifts):
        c = ws.cell(hdr, 3 + len(mines) + j, f"{u['host'].replace('_', ' ').title()}: {u['name'].replace('_', ' ')}")
        c.font = BOLD; c.alignment = Alignment(horizontal="right", wrap_text=True)
    r[0] += 1
    MROW = {}   # attribute -> row number on Assumptions, all mine columns

    def mine_ref(attr, j):
        return f"'Assumptions'!${_col(3 + j)}${MROW[attr]}"

    def typed_row(attr, label, values, fmt, cols=None):
        """Yellow inputs across the mine columns given (default: the guidance mines)."""
        ws.cell(r[0], 1, label)
        cols = range(len(mines)) if cols is None else cols
        for j, v in zip(cols, values):
            c = ws.cell(r[0], 3 + j, v); c.fill = YELLOW; c.font = BLUEFONT
            if fmt:
                c.number_format = fmt
        MROW[attr] = r[0]
        r[0] += 1

    def formula_row(attr, label, fn, fmt, cols=None, bold=False):
        """``fn(j, colletter)`` returns the formula for mine column j."""
        c0 = ws.cell(r[0], 1, label)
        if bold:
            c0.font = BOLD
        cols = range(ncols_mines) if cols is None else cols
        for j in cols:
            c = ws.cell(r[0], 3 + j, fn(j, _col(3 + j))); c.number_format = fmt
            if bold:
                c.font = BOLD
        MROW[attr] = r[0]
        r[0] += 1

    guidance = [("gold_koz", "Gold (koz pa)", NUM1), ("copper_kt", "Copper (kt pa)", NUM1),
                ("aisc_aud_oz", f"AISC ({ccy}/oz, net of by-product credits, incl. royalties at guidance prices)", NUM0),
                ("sustaining_capex", f"Sustaining capital ({ccy} m pa, inside AISC)", NUM1),
                ("major_mine_development", f"Major mine development ({ccy} m pa)", NUM1),
                ("exploration", f"Exploration ({ccy} m pa)", NUM1),
                ("da_aud_oz", f"Book D&A ({ccy}/oz, the tax-depreciation proxy)", NUM0),
                ("major_project_capex", f"Major project capital, FY27 ({ccy} m)", NUM1)]
    for key, label, fmt in guidance:
        typed_row(f"{key}_lo", label + " — low", [g["mines"][m][key]["low"] for m in mines], fmt)
        typed_row(f"{key}_hi", label + " — high", [g["mines"][m][key]["high"] for m in mines], fmt)
    typed_row("life_end", "Stated life of mine, last fiscal year", [fin["stated_mine_life_end_fy"][m] for m in mines], "0")
    typed_row("res_g", "Ore Reserves, gold (Moz)", [res[m]["gold_moz"] for m in mines], NUM2)
    typed_row("res_c", "Ore Reserves, copper (kt)", [res[m]["copper_kt"] for m in mines], NUM0)
    made = fin["production_since_reserve_date"]
    typed_row("made_g", f"Produced since the reserve date, gold (koz; {made['window']})", [made["gold_koz"].get(m, 0.0) for m in mines], NUM1)
    typed_row("made_c", "Produced since the reserve date, copper (kt)", [made["copper_kt"].get(m, 0.0) for m in mines], NUM1)
    sd = [stream["gold_delivered_fraction_of_attributable"] if stream is not None and m == stream.get("mine", "northparkes") else 0.0 for m in mines]
    sc = [stream["cash_received_fraction_of_spot"] if stream is not None and m == stream.get("mine", "northparkes") else 0.0 for m in mines]
    typed_row("stream_sd", "Stream: share of gold delivered under the stream", sd, PCT2)
    typed_row("stream_sc", "Stream: cash received as a share of spot on delivered ounces", sc, PCT2)
    # uplift inputs (their own columns)
    U = range(len(mines), ncols_mines)
    host_idx = {u["name"]: mines.index(u["host"]) for u in uplifts}
    typed_row("u_start_fy", "Uplift: first fiscal year", [u["start_fy"] for u in uplifts], "0", U)
    typed_row("u_years", "Uplift: life in years (a number, or host_life)", [u["years"] for u in uplifts], NUM1, U)
    typed_row("u_gold_abs", "Uplift: gold (koz pa), when stated outright", [u.get("gold_koz_pa") for u in uplifts], NUM1, U)
    typed_row("u_cu_abs", "Uplift: copper (kt pa), when stated outright", [u.get("copper_kt_pa") for u in uplifts], NUM1, U)
    typed_row("u_frac_g", "Uplift: gold as a fraction of the host, when stated that way", [(u.get("fraction_of_host") or {}).get("gold") for u in uplifts], PCT2, U)
    typed_row("u_frac_c", "Uplift: copper as a fraction of the host, when stated that way", [(u.get("fraction_of_host") or {}).get("copper") for u in uplifts], PCT2, U)
    typed_row("u_basis", "Uplift: cost basis (host_unit_cost or incremental_recovery)", [u["cost_basis"] for u in uplifts], None, U)
    typed_row("u_host", "Uplift: host mine column index (1 = first mine column)", [host_idx[u["name"]] + 1 for u in uplifts], "0", U)
    r[0] += 1

    sub("Derived per mine (formulas; the run-off sheets read these rows)")
    MR = range(len(mines))

    def host(attr, j):
        """INDEX across the guidance-mine columns of ``attr`` at the uplift's host index."""
        c0, c1 = _col(3), _col(2 + len(mines))
        return f"INDEX($%s$%d:$%s$%d,1,%s%d)" % (c0, MROW[attr], c1, MROW[attr], _col(3 + j), MROW["u_host"])

    for key, label, fmt in guidance:
        if key == "major_project_capex":
            continue
        formula_row(key, "→ " + label + " — mid", lambda j, L, k=key: f"=({L}{MROW[k + '_lo']}+{L}{MROW[k + '_hi']})/2", fmt, MR)
    formula_row("major_project_capex", f"→ Major project capital, FY27 ({ccy} m) — mid",
                lambda j, L: f"=({L}{MROW['major_project_capex_lo']}+{L}{MROW['major_project_capex_hi']})/2", NUM1, MR)
    for j in U:
        ws.cell(MROW["major_project_capex"], 3 + j, "=0").number_format = NUM1  # uplift project capital sits in the host's (section 5.3)
    formula_row("start", "→ Start period (0 = producing now; uplifts: first FY less the first fiscal year)",
                lambda j, L: "=0" if j < len(mines) else f"={L}{MROW['u_start_fy']}-{R['fy1']}", "0")
    # inventory on each basis, guidance mines
    formula_row("inv_g_life", "→ Inventory on stated life, gold (koz) = rate x (stub + life end - first FY)",
                lambda j, L: f"={L}{MROW['gold_koz']}*({R['f0']}+{L}{MROW['life_end']}-{R['fy1']})", NUM1, MR)
    formula_row("inv_c_life", "→ Inventory on stated life, copper (kt)",
                lambda j, L: f"={L}{MROW['copper_kt']}*({R['f0']}+{L}{MROW['life_end']}-{R['fy1']})", NUM1, MR)
    formula_row("inv_g_res", "→ Inventory on Ore Reserves, gold (koz) = Moz x 1,000 x recovery - produced since",
                lambda j, L: f"={L}{MROW['res_g']}*{_KOZ_PER_MOZ:g}*{R['rec']}-{L}{MROW['made_g']}", NUM1, MR)
    formula_row("inv_c_res", "→ Inventory on Ore Reserves, copper (kt) = reserve x recovery - produced since",
                lambda j, L: f"={L}{MROW['res_c']}*{R['rec']}-{L}{MROW['made_c']}", NUM1, MR)
    formula_row("inv_g_ext", "→ Inventory on stated life + extension, gold (koz)",
                lambda j, L: f"={L}{MROW['gold_koz']}*({R['f0']}+{L}{MROW['life_end']}-{R['fy1']}+{R['ext']})", NUM1, MR)
    formula_row("inv_c_ext", "→ Inventory on stated life + extension, copper (kt)",
                lambda j, L: f"={L}{MROW['copper_kt']}*({R['f0']}+{L}{MROW['life_end']}-{R['fy1']}+{R['ext']})", NUM1, MR)
    # basis 1 stated_life, 2 reserves, 3 life_extension (BASIS_ORDER)
    formula_row("inv_g", "→ Inventory on the SELECTED basis, gold (koz)",
                lambda j, L: f"=CHOOSE({R['basis']},{L}{MROW['inv_g_life']},{L}{MROW['inv_g_res']},{L}{MROW['inv_g_ext']})", NUM1, MR, bold=True)
    formula_row("inv_c", "→ Inventory on the SELECTED basis, copper (kt)",
                lambda j, L: f"=CHOOSE({R['basis']},{L}{MROW['inv_c_life']},{L}{MROW['inv_c_res']},{L}{MROW['inv_c_ext']})", NUM1, MR, bold=True)
    # uplift derivations (host-driven), written into the uplift columns of the same attribute rows
    for j in U:
        L = _col(3 + j)
        hg, hc = host("gold_koz", j), host("copper_kt", j)
        ws.cell(MROW["gold_koz"], 3 + j, f"=IF(ISNUMBER({L}{MROW['u_frac_g']}),{hg}*{L}{MROW['u_frac_g']},{L}{MROW['u_gold_abs']})").number_format = NUM1
        ws.cell(MROW["copper_kt"], 3 + j, f"=IF(ISNUMBER({L}{MROW['u_frac_c']}),{hc}*{L}{MROW['u_frac_c']},{L}{MROW['u_cu_abs']})").number_format = NUM1
        hig, hic = host("inv_g", j), host("inv_c", j)
        host_years = f"MIN(IF({hg}>0,{hig}/{hg},{_BIG:g}),IF({hc}>0,{hic}/{hc},{_BIG:g}))"
        ws.cell(MROW["inv_g_life"], 3 + j, f"=IF(ISNUMBER({L}{MROW['u_years']}),{L}{MROW['u_years']},{host_years}-({R['f0']}+{L}{MROW['start']}-1))").number_format = NUM2
        ws.cell(MROW["inv_c_life"], 3 + j, "(uplift life in years: row above)").font = NOTE
        ws.cell(MROW["inv_g"], 3 + j, f"={L}{MROW['gold_koz']}*{L}{MROW['inv_g_life']}").number_format = NUM1
        ws.cell(MROW["inv_c"], 3 + j, f"={L}{MROW['copper_kt']}*{L}{MROW['inv_g_life']}").number_format = NUM1
        rev_u = f"({L}{MROW['gold_koz']}*{R['pb_g']}+{L}{MROW['copper_kt']}*{R['pb_c']})"
        rev_h = f"({hg}*{R['pb_g']}+{hc}*{R['pb_c']})"
        gross_h = f"({host('aisc_aud_oz', j)}*{hg}+{hc}*{R['pb_c']})"
        is_host = f'({L}{MROW["u_basis"]}="host_unit_cost")'
        gross_u = f"IF({is_host},{gross_h}*{rev_u}/{rev_h},{R['roy']}*{rev_u})"
        ws.cell(MROW["aisc_aud_oz"], 3 + j, f"=IF({L}{MROW['gold_koz']}>0,({gross_u}-{L}{MROW['copper_kt']}*{R['pb_c']})/{L}{MROW['gold_koz']},0)").number_format = NUM0
        ws.cell(MROW["sustaining_capex"], 3 + j, f"=IF({is_host},{host('sustaining_capex', j)}*{rev_u}/{rev_h},0)").number_format = NUM2
        ws.cell(MROW["major_mine_development"], 3 + j, f"=IF({is_host},{host('major_mine_development', j)}*{rev_u}/{rev_h},0)").number_format = NUM2
        ws.cell(MROW["exploration"], 3 + j, "=0").number_format = NUM1
        ws.cell(MROW["da_aud_oz"], 3 + j, f"={host('da_aud_oz', j)}").number_format = NUM0
        ws.cell(MROW["stream_sd"], 3 + j, f"={host('stream_sd', j)}").number_format = PCT2
        ws.cell(MROW["stream_sc"], 3 + j, f"={host('stream_sc', j)}").number_format = PCT2
    ws.cell(MROW["inv_g_life"], 1, ws.cell(MROW["inv_g_life"], 1).value + "; uplift columns: life in years")
    formula_row("gross", f"→ Gross cash cost at guidance prices ({ccy} m pa) = AISC x gold + copper x guidance copper price (credit added back)",
                lambda j, L: f"=({L}{MROW['aisc_aud_oz']}*{L}{MROW['gold_koz']}+{L}{MROW['copper_kt']}*{R['pb_c']})/{_THOUSANDS:g}", NUM1)
    formula_row("g_roy", f"→ Royalty inside guidance AISC ({ccy} m pa) = rate x guidance revenue",
                lambda j, L: f"={R['roy']}*({L}{MROW['gold_koz']}*{R['pb_g']}+{L}{MROW['copper_kt']}*{R['pb_c']})/{_THOUSANDS:g}", NUM1)
    formula_row("haircut", "→ Stream haircut on gold revenue = delivered share x (1 - cash share of spot)",
                lambda j, L: f"={L}{MROW['stream_sd']}*(1-{L}{MROW['stream_sc']})", PCT2)
    r[0] += 1

    # ---- scenario series anchors ----
    sub("Scenario series anchors (year, value), the scenario files' own points; interpolated linearly between them and held flat outside (D-37)")
    SROW = {}  # (sid, series) -> (row of years, row of values, n)
    ESC = {}
    EQ = {}
    for sid in SIDS:
        si = by_sid[sid]
        ws.cell(r[0], 1, NAMES[sid]).font = BOLD
        r[0] += 1
        for skey, label in (("gold", f"Gold, real 2026 US$/oz ({series['gold']})"),
                            ("copper", f"Copper, real 2026 US$/t ({series['copper']})"),
                            ("fx", f"AUD/USD ({series['fx']})"),
                            ("cpi", f"CPI, % pa ({series['cpi']})")):
            pts = _series_anchors(si, series[skey])
            ws.cell(r[0], 1, f"  {label} — anchor years")
            for k, (y, v) in enumerate(pts):
                c = ws.cell(r[0], 2 + k, y); c.fill = YELLOW; c.font = BLUEFONT; c.number_format = "0"
            ws.cell(r[0] + 1, 1, f"  {label} — anchor values")
            for k, (y, v) in enumerate(pts):
                c = ws.cell(r[0] + 1, 2 + k, v); c.fill = YELLOW; c.font = BLUEFONT
                c.number_format = NUM2 if skey in ("fx", "cpi") else NUM0
            SROW[(sid, skey)] = (r[0], r[0] + 1, len(pts))
            r[0] += 2
        ws.cell(r[0], 1, "  Real cost escalation above CPI, % pa (impact matrix mid; from year 1 to the year before the final phase)")
        c = ws.cell(r[0], 2, _escalation_mid(si, sid)); c.fill = YELLOW; c.font = BLUEFONT; c.number_format = NUM2
        ESC[sid] = f"'Assumptions'!$B${r[0]}"
        r[0] += 1
        ws.cell(r[0], 1, "  Year the scenario's final phase starts (time profile)")
        c = ws.cell(r[0], 2, int(scenario_equilibrium_year(si))); c.fill = YELLOW; c.font = BLUEFONT; c.number_format = "0"
        EQ[sid] = f"'Assumptions'!$B${r[0]}"
        r[0] += 1
    ws.column_dimensions["A"].width = 78
    for j in range(2, 4 + ncols_mines):
        ws.column_dimensions[_col(j)].width = 16

    # ======================= Prices =======================
    wp = wb.create_sheet("Prices")
    wp.cell(1, 1, f"Prices and cost indices by period (period 0 = stub; the stub reads scenario year 1). {ccy} prices = real US$ x CPI index / AUD-USD.").font = HDR
    wp.cell(2, 1, "Interpolation: linear between the anchor years on Assumptions, held flat before the first and after the last (D-37). The D-104 multiplier reads the real A$ gold price against period 0, with the lag and cap.").font = NOTE
    PR = {}   # (sid, row key) -> row number on Prices
    pr = 4
    FIRST = 3  # first period column (C)

    def interp(sid, skey, ycell):
        yr, vr, n = SROW[(sid, skey)]
        Y = f"'Assumptions'!$B${yr}:${_col(1 + n)}${yr}"
        V = f"'Assumptions'!$B${vr}:${_col(1 + n)}${vr}"
        m = f"MATCH({ycell},{Y},1)"
        first_y, last_y = f"'Assumptions'!$B${yr}", f"'Assumptions'!${_col(1 + n)}${yr}"
        first_v, last_v = f"'Assumptions'!$B${vr}", f"'Assumptions'!${_col(1 + n)}${vr}"
        lin = (f"INDEX({V},1,{m})+(INDEX({V},1,{m}+1)-INDEX({V},1,{m}))*({ycell}-INDEX({Y},1,{m}))"
               f"/(INDEX({Y},1,{m}+1)-INDEX({Y},1,{m}))")
        return f"IF({ycell}<={first_y},{first_v},IF({ycell}>={last_y},{last_v},{lin}))"

    for sid in SIDS:
        wp.cell(pr, 1, NAMES[sid]).font = SUB
        pr += 1
        rows = ["t", "year", "cpi", "cpi_idx", "esc", "cost_idx", "fx", "gold_usd", "cu_usd", "gold_aud", "cu_aud", "gold_aud_real", "link"]
        labels = {"t": "Period t", "year": "Scenario year read (max(t, 1))", "cpi": "CPI, % pa (year t)",
                  "cpi_idx": "CPI index (1.0 in the stub)", "esc": "Real cost escalation, % pa",
                  "cost_idx": "Cost index (CPI plus escalator)", "fx": "AUD/USD",
                  "gold_usd": "Gold, real US$/oz", "cu_usd": "Copper, real US$/t",
                  "gold_aud": f"Gold, nominal {ccy}/oz", "cu_aud": f"Copper, nominal {ccy}/t",
                  "gold_aud_real": f"Gold, real {ccy}/oz (nominal / CPI index)", "link": "D-104 cost multiplier"}
        for k in rows:
            PR[(sid, k)] = pr
            wp.cell(pr, 1, labels[k])
            pr += 1
        pr += 1
        for t in range(NP):
            C = _col(FIRST + t)
            P = _col(FIRST + t - 1)
            rt = PR[(sid, "t")]
            wp.cell(rt, FIRST + t, f"=COLUMN()-{FIRST}").number_format = "0"
            wp.cell(PR[(sid, "year")], FIRST + t, f"=MAX({C}{rt},1)").number_format = "0"
            wp.cell(PR[(sid, "cpi")], FIRST + t, f"={interp(sid, 'cpi', f'{C}{rt}')}" if t > 0 else "=0").number_format = NUM2
            wp.cell(PR[(sid, "cpi_idx")], FIRST + t, "=1" if t == 0 else f"={P}{PR[(sid, 'cpi_idx')]}*(1+{C}{PR[(sid, 'cpi')]}/100)").number_format = NUM4  # ssot-allow: percent to decimal
            wp.cell(PR[(sid, "esc")], FIRST + t, f"=IF(AND({C}{rt}>=1,{C}{rt}<{EQ[sid]}),{ESC[sid]},0)").number_format = NUM2
            wp.cell(PR[(sid, "cost_idx")], FIRST + t, "=1" if t == 0 else f"={P}{PR[(sid, 'cost_idx')]}*(1+{C}{PR[(sid, 'cpi')]}/100+{C}{PR[(sid, 'esc')]}/100)").number_format = NUM4  # ssot-allow: percent to decimal
            yc = f"{C}{PR[(sid, 'year')]}"
            wp.cell(PR[(sid, "fx")], FIRST + t, f"={interp(sid, 'fx', yc)}").number_format = NUM4
            wp.cell(PR[(sid, "gold_usd")], FIRST + t, f"={interp(sid, 'gold', yc)}").number_format = NUM0
            wp.cell(PR[(sid, "cu_usd")], FIRST + t, f"={interp(sid, 'copper', yc)}").number_format = NUM0
            wp.cell(PR[(sid, "gold_aud")], FIRST + t, f"={C}{PR[(sid, 'gold_usd')]}*{C}{PR[(sid, 'cpi_idx')]}/{C}{PR[(sid, 'fx')]}").number_format = NUM0
            wp.cell(PR[(sid, "cu_aud")], FIRST + t, f"={C}{PR[(sid, 'cu_usd')]}*{C}{PR[(sid, 'cpi_idx')]}/{C}{PR[(sid, 'fx')]}").number_format = NUM0
            wp.cell(PR[(sid, "gold_aud_real")], FIRST + t, f"={C}{PR[(sid, 'gold_aud')]}/{C}{PR[(sid, 'cpi_idx')]}").number_format = NUM0
            # link: period t reads the real price `lag` periods earlier (period 0's own level before the lag can read anything)
            ref = f"${_col(FIRST)}${PR[(sid, 'gold_aud_real')]}"
            lagged = f"IF({C}{rt}-{R['link_lag_years']}<0,{ref},INDEX(${_col(FIRST)}${PR[(sid, 'gold_aud_real')]}:${_col(FIRST + NP - 1)}${PR[(sid, 'gold_aud_real')]},1,{C}{rt}-{R['link_lag_years']}+1))"
            raw = f"MIN(1+{R['link_elasticity_per_10pct']}/10*({lagged}/{ref}-1),1+{R['link_cap_above_base_pct']}/100)"  # ssot-allow: per-10% elasticity to per-unit move, as the translator does
            wp.cell(PR[(sid, "link")], FIRST + t, f"=IF({R['link_on']}=0,1,IF({R['link_symmetric']}=1,{raw},MAX({raw},1)))").number_format = NUM4
    wp.column_dimensions["A"].width = 44
    for t in range(NP):
        wp.column_dimensions[_col(FIRST + t)].width = 11

    # ======================= Run-off sheets =======================
    VAL = {}  # (sid, key) -> "'Run-off X'!$B$r"
    for sid in SIDS:
        sh = f"Run-off {SHORT[sid]}"
        wr = wb.create_sheet(sh)
        wr.cell(1, 1, f"{NAMES[sid]} — mine-by-mine depletion run-off ({ccy} m unless stated). Period 0 is the stub; columns run to the longest life on any inventory basis.").font = HDR
        wr.cell(2, 1, "Each mine produces at its guidance rate until its first metal runs out (ore-driven); costs are guidance AISC with the by-product credit added back and the royalty stripped, escalated by the cost index; non-royalty cash cost and sustaining capital carry the D-104 multiplier.").font = NOTE
        rr = 4
        ROW = {}

        def rowf(key, label, fn, fmt=NUM1, bold=False):
            """Period row: ``fn(t, C, P)`` -> formula for column t (C its letter, P the prior column's)."""
            nonlocal rr
            c0 = wr.cell(rr, 1, label)
            if bold:
                c0.font = BOLD
            for t in range(NP):
                f = fn(t, _col(FIRST + t), _col(FIRST + t - 1))
                c = wr.cell(rr, FIRST + t, f); c.number_format = fmt
                if bold:
                    c.font = BOLD
            ROW[key] = rr
            rr += 1

        def pref(key, C):
            return f"{C}${ROW[key]}"

        wr.cell(rr, 1, "Period and prices (from Prices and Assumptions)").font = SUB; rr += 1
        rowf("t", "Period t", lambda t, C, P: f"=COLUMN()-{FIRST}", "0")
        rowf("fy", "Fiscal year", lambda t, C, P: f"={R['fy1']}+{C}${ROW['t']}", "0")
        rowf("frac", "Year fraction (stub in period 0)", lambda t, C, P: f"=IF({C}${ROW['t']}=0,{R['f0']},1)", NUM4)
        rowf("ci", "Cost index", lambda t, C, P: f"='Prices'!{C}${PR[(sid, 'cost_idx')]}", NUM4)
        rowf("pg", f"Gold price ({ccy}/oz, nominal)", lambda t, C, P: f"='Prices'!{C}${PR[(sid, 'gold_aud')]}", NUM0)
        rowf("pc", f"Copper price ({ccy}/t, nominal)", lambda t, C, P: f"='Prices'!{C}${PR[(sid, 'cu_aud')]}", NUM0)
        rowf("link", "D-104 cost multiplier", lambda t, C, P: f"='Prices'!{C}${PR[(sid, 'link')]}", NUM4)
        rr += 1

        mine_rows = []
        for j in range(ncols_mines):
            name = (mines[j].replace("_", " ").title() if j < len(mines)
                    else f"{uplifts[j - len(mines)]['host'].replace('_', ' ').title()}: {uplifts[j - len(mines)]['name'].replace('_', ' ')}")
            wr.cell(rr, 1, name).font = SUB; rr += 1
            M = {k: mine_ref(k, j) for k in ("gold_koz", "copper_kt", "inv_g", "inv_c", "aisc_aud_oz", "sustaining_capex",
                                            "major_mine_development", "exploration", "da_aud_oz", "major_project_capex",
                                            "start", "gross", "g_roy", "haircut")}
            k = f"m{j}:"
            for i, nm in enumerate(("rem_g", "rem_c", "scale", "gold", "cu", "rev", "opex", "expl", "capex_s", "capex_p", "da")):
                ROW[k + nm] = rr + i   # pre-assigned so a row may reference one written after it
            rowf(k + "rem_g", "  Remaining gold at the start of the period (koz)",
                 lambda t, C, P, M=M, k=k: f"={M['inv_g']}" if t == 0 else f"={P}${ROW[k + 'rem_g']}-{P}${ROW[k + 'gold']}", NUM1)
            rowf(k + "rem_c", "  Remaining copper at the start of the period (kt)",
                 lambda t, C, P, M=M, k=k: f"={M['inv_c']}" if t == 0 else f"={P}${ROW[k + 'rem_c']}-{P}${ROW[k + 'cu']}", NUM1)
            rowf(k + "scale", "  Production scale (year fraction, cut when a metal runs out; 0 before the start period)",
                 lambda t, C, P, M=M, k=k: (f"=IF({C}${ROW['t']}<{M['start']},0,"
                                            f"IF(MIN({C}${ROW['frac']},IF({M['gold_koz']}>0,MAX({C}${ROW[k + 'rem_g']},0)/{M['gold_koz']},{_BIG:g}),"
                                            f"IF({M['copper_kt']}>0,MAX({C}${ROW[k + 'rem_c']},0)/{M['copper_kt']},{_BIG:g}))<={_EXHAUSTED:g},0,"
                                            f"MIN({C}${ROW['frac']},IF({M['gold_koz']}>0,MAX({C}${ROW[k + 'rem_g']},0)/{M['gold_koz']},{_BIG:g}),"
                                            f"IF({M['copper_kt']}>0,MAX({C}${ROW[k + 'rem_c']},0)/{M['copper_kt']},{_BIG:g}))))"), NUM4)
            rowf(k + "gold", "  Gold sold (koz)", lambda t, C, P, M=M, k=k: f"={M['gold_koz']}*{C}${ROW[k + 'scale']}", NUM1)
            rowf(k + "cu", "  Copper sold (kt)", lambda t, C, P, M=M, k=k: f"={M['copper_kt']}*{C}${ROW[k + 'scale']}", NUM1)
            rowf(k + "rev", "  Revenue (gold net of the stream haircut, plus copper)",
                 lambda t, C, P, M=M, k=k: f"=({C}${ROW[k + 'gold']}*{C}${ROW['pg']}*(1-{M['haircut']})+{C}${ROW[k + 'cu']}*{C}${ROW['pc']})/{_THOUSANDS:g}", NUM1)
            rowf(k + "opex", "  Operating cost = (gross cash cost - guidance royalty - sustaining) x scale x cost index x D-104",
                 lambda t, C, P, M=M, k=k: f"=({M['gross']}-{M['g_roy']}-{M['sustaining_capex']})*{C}${ROW[k + 'scale']}*{C}${ROW['ci']}*{C}${ROW['link']}", NUM1)
            rowf(k + "expl", "  Exploration x scale x cost index",
                 lambda t, C, P, M=M, k=k: f"={M['exploration']}*{C}${ROW[k + 'scale']}*{C}${ROW['ci']}", NUM1)
            rowf(k + "capex_s", "  Sustaining (x D-104) and mine development x scale x cost index",
                 lambda t, C, P, M=M, k=k: f"=({M['sustaining_capex']}*{C}${ROW['link']}+{M['major_mine_development']})*{C}${ROW[k + 'scale']}*{C}${ROW['ci']}", NUM1)
            rowf(k + "capex_p", "  Project capital = FY27 mid x run-off fraction x year fraction x cost index (nil after the run-off)",
                 lambda t, C, P, M=M, k=k: (f"=IF({C}${ROW[k + 'scale']}>0,1,0)*{M['major_project_capex']}*{R['runoff' + str(t)]}*{C}${ROW['frac']}*{C}${ROW['ci']}" if t < len(runoff) else "=0"), NUM1)
            rowf(k + "da", "  D&A = book D&A per ounce x gold sold",
                 lambda t, C, P, M=M, k=k: f"={M['da_aud_oz']}*{C}${ROW[k + 'gold']}/{_THOUSANDS:g}", NUM1)
            mine_rows.append(k)
            rr += 1

        def sum_mines(attr, C):
            return "+".join(f"{C}${ROW[k + attr]}" for k in mine_rows)

        wr.cell(rr, 1, "Group").font = SUB; rr += 1
        rowf("prod", "Producing? (1 while any mine produces)", lambda t, C, P: f"=IF(({sum_mines('scale', C)})>0,1,0)", "0")
        rowf("gold", "Group gold sold (koz)", lambda t, C, P: f"={sum_mines('gold', C)}", NUM1)
        rowf("cu", "Group copper sold (kt)", lambda t, C, P: f"={sum_mines('cu', C)}", NUM1)
        rowf("rev", "Revenue", lambda t, C, P: f"={sum_mines('rev', C)}", NUM1, bold=True)
        rowf("roy", "Royalty = rate x revenue", lambda t, C, P: f"={R['roy']}*{C}${ROW['rev']}", NUM1)
        rowf("opex", "Operating cost", lambda t, C, P: f"={sum_mines('opex', C)}", NUM1)
        rowf("expl", "Exploration and corporate = mine exploration + (corporate per oz x gold + greenfields x year fraction) x cost index, while producing",
             lambda t, C, P: f"={sum_mines('expl', C)}+{C}${ROW['prod']}*({R['corp']}*{C}${ROW['gold']}/{_THOUSANDS:g}+{R['green']}*{C}${ROW['frac']})*{C}${ROW['ci']}", NUM1)
        rowf("ebitda", "EBITDA", lambda t, C, P: f"={C}${ROW['rev']}-{C}${ROW['roy']}-{C}${ROW['opex']}-{C}${ROW['expl']}", NUM1, bold=True)
        rowf("da", "D&A (tax-depreciation proxy)", lambda t, C, P: f"={sum_mines('da', C)}", NUM1)
        rowf("tax", "Tax = rate x max(EBITDA - D&A, 0)", lambda t, C, P: f"={R['tax']}*MAX({C}${ROW['ebitda']}-{C}${ROW['da']},0)", NUM1)
        rowf("capex", "Capital = sustaining, mine development and project capital",
             lambda t, C, P: f"={sum_mines('capex_s', C)}+{sum_mines('capex_p', C)}", NUM1)
        rowf("rev_ann", "Revenue, annualised (revenue / year fraction)", lambda t, C, P: f"={C}${ROW['rev']}/{C}${ROW['frac']}", NUM1)
        rowf("dwc", "Change in working capital = intensity x change in annualised revenue; the balance unwinds in the last producing period",
             lambda t, C, P: (f"=IF({C}${ROW['prod']}=1,"
                              + ("0" if t == 0 else f"{R['wci']}*({C}${ROW['rev_ann']}-{P}${ROW['rev_ann']})")
                              + f"+IF({(_col(FIRST + t + 1) + '$' + str(ROW['prod'])) if t + 1 < NP else '0'}=1,0,{R['wci']}*(0-{C}${ROW['rev_ann']})),0)"), NUM1)
        rowf("fcff", "FCFF = EBITDA - tax - capital - change in working capital",
             lambda t, C, P: f"={C}${ROW['ebitda']}-{C}${ROW['tax']}-{C}${ROW['capex']}-{C}${ROW['dwc']}", NUM1, bold=True)
        rowf("mid", "Mid-period time from the valuation date (years)",
             lambda t, C, P: f"=IF({C}${ROW['t']}=0,{R['f0']}/2,{R['f0']}+{C}${ROW['t']}-0.5)", NUM4)
        rowf("df", "Discount factor at the WACC", lambda t, C, P: f"=1/(1+{R['wacc']})^{C}${ROW['mid']}", NUM4)
        rowf("pv", "PV of FCFF", lambda t, C, P: f"={C}${ROW['fcff']}*{C}${ROW['df']}", NUM1, bold=True)
        rowf("explicit", "In the explicit period? (1 for the stub and the horizon years)", lambda t, C, P: f"=IF({C}${ROW['t']}<={R['H']},1,0)", "0")
        rr += 1

        wr.cell(rr, 1, "Valuation").font = SUB; rr += 1
        F, LST = _col(FIRST), _col(FIRST + NP - 1)

        def rng(key):
            return f"${F}${ROW[key]}:${LST}${ROW[key]}"

        def scalar(key, label, formula, fmt=NUM1, bold=False):
            nonlocal rr
            c0 = wr.cell(rr, 1, label)
            c = wr.cell(rr, 2, formula); c.number_format = fmt
            if bold:
                c0.font = BOLD; c.font = BOLD
            VAL[(sid, key)] = f"'{sh}'!$B${rr}"
            ROW[key] = rr
            rr += 1

        scalar("pv_explicit", "PV of explicit FCFF (stub and the horizon years)", f"=SUMPRODUCT({rng('pv')},{rng('explicit')})")
        scalar("pv_terminal", "PV of life-of-mine cash flows after the explicit period (the depletion run-off, D-97)", f"=SUMPRODUCT({rng('pv')},1-{rng('explicit')})")
        scalar("terminal_value", "Run-off value at the end of the explicit period (PV run-off x (1 + WACC)^(stub + horizon))",
               f"={VAL[(sid, 'pv_terminal')]}*(1+{R['wacc']})^({R['f0']}+{R['H']})")
        scalar("enterprise_value", "Enterprise value", f"={VAL[(sid, 'pv_explicit')]}+{VAL[(sid, 'pv_terminal')]}", bold=True)
        scalar("net_debt", "Less: net debt at valuation", f"={R['net_debt']}")
        scalar("claims", "Less: rehabilitation and other claims", f"={R['claims']}")
        scalar("leases", "Less: leases", f"={R['leases']}")
        scalar("equity_value", "Equity value", f"={VAL[(sid, 'enterprise_value')]}-{VAL[(sid, 'net_debt')]}-{VAL[(sid, 'claims')]}-{VAL[(sid, 'leases')]}", bold=True)
        scalar("shares", "Shares outstanding (m)", f"={R['shares']}", NUM1 + "00")
        scalar("value_per_share", f"Value per share ({ccy})", f"={VAL[(sid, 'equity_value')]}/{VAL[(sid, 'shares')]}", NUM2, bold=True)
        scalar("vs_market", "vs market price", f"={VAL[(sid, 'value_per_share')]}/{R['mkt']}-1", PCT2)
        scalar("terminal_share_of_ev", "Run-off share of enterprise value", f"={VAL[(sid, 'pv_terminal')]}/{VAL[(sid, 'enterprise_value')]}", PCT2)
        scalar("life_years", "Life, years from the valuation date to the last closure (stub + producing periods - 1; production is contiguous from period 0, as the engine's loop assumes)",
               f"=IF(SUM({rng('prod')})>1,{R['f0']}+SUM({rng('prod')})-1,{R['f0']})", NUM2)
        scalar("periods", "Producing periods (the engine's n)", f"=SUM({rng('prod')})", "0")
        wr.column_dimensions["A"].width = 70
        wr.column_dimensions["B"].width = 14
        for t in range(NP):
            wr.column_dimensions[_col(FIRST + t)].width = 10.5
        wr.freeze_panes = wr.cell(5, FIRST)

    # ======================= Scenarios =======================
    wsx = wb.create_sheet("Scenarios")
    wsx.cell(1, 1, f"Scenario summary — value per share ({ccy}); links to the run-off sheets").font = HDR
    hr = 3
    for j, t in enumerate(["Scenario", f"{ccy}/share", f"EV ({ccy} m)", "vs market", "Run-off share of EV", "Life (years)", "PV explicit", "PV run-off"]):
        c = wsx.cell(hr, 1 + j, t); c.font = BOLD; c.alignment = Alignment(wrap_text=True)
    rr2 = hr + 1
    for sid in SIDS:
        wsx.cell(rr2, 1, NAMES[sid])
        wsx.cell(rr2, 2, f"={VAL[(sid, 'value_per_share')]}").number_format = NUM2
        wsx.cell(rr2, 3, f"={VAL[(sid, 'enterprise_value')]}").number_format = NUM0
        wsx.cell(rr2, 4, f"={VAL[(sid, 'vs_market')]}").number_format = PCT2
        wsx.cell(rr2, 5, f"={VAL[(sid, 'terminal_share_of_ev')]}").number_format = PCT2
        wsx.cell(rr2, 6, f"={VAL[(sid, 'life_years')]}").number_format = NUM2
        wsx.cell(rr2, 7, f"={VAL[(sid, 'pv_explicit')]}").number_format = NUM0
        wsx.cell(rr2, 8, f"={VAL[(sid, 'pv_terminal')]}").number_format = NUM0
        rr2 += 1
    rr2 += 1
    wsx.cell(rr2, 1, f"Market price ({ccy})"); wsx.cell(rr2, 2, f"={R['mkt']}").number_format = NUM2; rr2 += 1
    wsx.cell(rr2, 1, "WACC"); wsx.cell(rr2, 2, f"={R['wacc']}").number_format = PCT2; rr2 += 1
    wsx.cell(rr2, 1, "Inventory basis selected (Assumptions)"); wsx.cell(rr2, 2, f"={R['basis']}").number_format = "0"; rr2 += 1
    wsx.cell(rr2 + 1, 1, "The terminal is the depletion run-off, not a growth perpetuity (D-97). Change the inventory basis, the WACC inputs, the price anchors or the D-104 rows on Assumptions and every sheet follows.").font = NOTE
    wsx.column_dimensions["A"].width = 30
    for cc in "BCDEFGH":
        wsx.column_dimensions[cc].width = 15

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    sys.path.insert(0, str(_ROOT / "scripts"))
    from size_resource_valuation import resource_companies
    for cid, arch in resource_companies():
        data = build_resource_workbook_bytes(cid, arch)
        out = _ROOT / "ui_prototypes" / "_generator" / f"_{cid}_full.xlsx"
        out.write_bytes(data)
        print("wrote", out, len(data), "bytes")

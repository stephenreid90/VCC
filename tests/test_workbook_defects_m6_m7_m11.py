"""Pins for three workbook defects from the August review (tracker items M6, M7, M11).

M6  DNL's two DCF views must carry the same EV and equity bridge, and the DCF-table
    tie check must compare to the production engine, not to a back-solve of itself.
M7  No generated file may name another company in the beta workbench label.
M11 Every company's Assumptions sheet carries the market price (yellow input) and
    the market reference date where the data has one; no other cell hard-codes the price.

The JS arithmetic is run in node when node is installed (skipped otherwise).
"""

from __future__ import annotations

import io
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
GEN = ROOT / "ui_prototypes" / "_generator"
sys.path.insert(0, str(GEN))
sys.path.insert(0, str(ROOT / "src"))

import engine_workbook as ew  # noqa: E402
import resource_workbook as rw  # noqa: E402

sys.path.insert(0, str(ROOT / "scripts"))
from size_resource_valuation import resource_companies  # noqa: E402

TICKERS = {"dnl": "DNL", "wbc": "WBC", "csl": "CSL"}
BUILDERS = {
    "dnl": ew.build_dnl_workbook_bytes,
    "wbc": ew.build_wbc_workbook_bytes,
    "csl": ew.build_csl_workbook_bytes,
}
ARCH = {"dnl": "industrial_explosives", "wbc": "australian_major_banks", "csl": "biopharmaceuticals"}
# Resource companies (gap 9, 20 Oct 2026) are held to M7 and M11 like the other three.
for _cid, _arch in resource_companies():
    TICKERS[_cid] = _cid.upper()
    BUILDERS[_cid] = (lambda c=_cid, a=_arch: rw.build_resource_workbook_bytes(c, a))
    ARCH[_cid] = _arch


def _market_reference(inp):
    """(price, date): the one reader for the three original companies, the financials
    ``market_data`` block for a resource company (its translator reads it from there)."""
    craw = inp["company_raw"]
    if (craw.get("normalised_baseline") or {}).get("engine") == "resource":
        md = inp["financials"]["market_data"]
        return md["price_aud"], md.get("price_date")
    return ew._market_reference(craw)
GEN_UI = (GEN / "gen_ui.py").read_text(encoding="utf-8")


@pytest.fixture(scope="module")
def books():
    return {c: load_workbook(io.BytesIO(b())) for c, b in BUILDERS.items()}


def _company_inputs(company: str):
    from vcc_valuations.translator import load_inputs
    return load_inputs(ROOT, "muddle_through", ARCH[company], company)


# ---------------------------------------------------------------- M7

@pytest.mark.parametrize("company", list(TICKERS))
def test_m7_no_other_company_named_in_workbook_text(books, company):
    others = [t for k, t in TICKERS.items() if k != company]
    pat = re.compile(r"\b(" + "|".join(others) + r")\b")
    hits = []
    for ws in books[company]:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith("=") and pat.search(c.value):
                    hits.append(f"{ws.title}!{c.coordinate}: {c.value[:80]}")
    assert not hits, hits


def test_m7_dnl_beta_sheet_names_dnl(books):
    ws = books["dnl"]["Comparables & beta"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str))
    assert "DNL D/E (for relever)" in text and "Relevered to DNL" in text


def test_m7_scaffold_relever_label_takes_the_company_from_config():
    assert "re-lever at DNL" not in GEN_UI
    assert "DNL’s gearing" not in GEN_UI
    assert "re-lever at '+CFG.companyShort+' gearing" in GEN_UI
    assert "at '+CFG.companyShort+'’s gearing" in GEN_UI


@pytest.mark.parametrize("company", list(TICKERS))
def test_m7_generated_prototype_has_no_leak(company):
    path = ROOT / "ui_prototypes" / f"{company}_scenario_interface.html"
    if not path.exists():
        pytest.skip("prototype not generated (gitignored build output)")
    text = path.read_text(encoding="utf-8")
    assert "DNL gearing" not in text and "DNL’s gearing" not in text


# ---------------------------------------------------------------- M11

@pytest.mark.parametrize("company", list(TICKERS))
def test_m11_market_price_is_a_yellow_input_on_assumptions(books, company):
    price, date = _market_reference(_company_inputs(company))
    ws = books[company]["Assumptions"]
    rows = [r for r in ws.iter_rows(min_col=1, max_col=2)
            if isinstance(r[0].value, str) and r[0].value.lower().startswith("market")
            and "price" in r[0].value.lower()]
    assert len(rows) == 1, [r[0].value for r in rows]
    cell = rows[0][1]
    assert cell.value == pytest.approx(price)
    assert cell.fill.fgColor.rgb == ew.YELLOW.fgColor.rgb
    assert cell.font.color is not None and cell.font.color.rgb == ew.BLUEFONT.color.rgb
    if date:
        drow = [r for r in ws.iter_rows(min_col=1, max_col=2)
                if r[0].value == "Market reference date"]
        assert len(drow) == 1 and drow[0][1].value is not None
        assert drow[0][1].fill.fgColor.rgb == ew.YELLOW.fgColor.rgb


@pytest.mark.parametrize("company", list(TICKERS))
def test_m11_no_other_cell_hard_codes_the_price(books, company):
    price, _ = _market_reference(_company_inputs(company))
    lit = re.compile(r"(?<![\d.])" + re.escape(repr(float(price)).rstrip("0").rstrip(".")) + r"(?![\d])")
    bad = []
    for ws in books[company]:
        if ws.title == "Assumptions":
            continue
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, (int, float)) and abs(v - price) < 1e-9:
                    bad.append(f"{ws.title}!{c.coordinate} value")
                elif isinstance(v, str) and v.startswith("=") and lit.search(v):
                    bad.append(f"{ws.title}!{c.coordinate} formula {v[:60]}")
    assert not bad, bad


@pytest.mark.parametrize("company", ["wbc", "csl"])
def test_m11_multiples_price_cell_is_a_link_to_assumptions(books, company):
    ws = books[company]["Multiples"]
    cells = [r[1] for r in ws.iter_rows(min_col=1, max_col=2)
             if isinstance(r[0].value, str) and r[0].value.lower().startswith("market price")]
    assert cells, "no market price row on Multiples"
    for c in cells:
        assert isinstance(c.value, str) and c.value.startswith("=") and "'Assumptions'!" in c.value


def test_m11_wbc_scenarios_vs_market_links_to_assumptions(books):
    ws = books["wbc"]["Scenarios"]
    forms = [c.value for row in ws.iter_rows(min_col=4, max_col=4) for c in row
             if isinstance(c.value, str) and c.value.startswith("=")]
    assert forms and all("'Assumptions'!" in f for f in forms), forms


# ---------------------------------------------------------------- M6

def _js_funcs():
    f1 = re.search(r"function dcfBridgeEv\(per\)\{.*?\}\n", GEN_UI).group(0)
    f2 = re.search(r"function dcfEngineTie\(name, ev\)\{.*?\}\n", GEN_UI).group(0)
    return f1 + f2


def _engine_pack():
    """Per-scenario engine EV and the bridge lines, straight from committed code."""
    from vcc_valuations.translator import (
        load_inputs, build_engine_inputs_from_data, equity_bridge_from_data)
    from vcc_valuations.dcf.fcf_engine import FcfEngine
    out, bridge = {}, None
    for sid, nm in ew.SCEN:
        inp = load_inputs(ROOT, sid, "industrial_explosives", "dnl")
        res = FcfEngine().run(build_engine_inputs_from_data(inp, sid))
        out[nm] = (res.enterprise_value, res.value_per_share)
        if sid == ew.CENTRAL:
            eb = equity_bridge_from_data(inp, sid)
            bridge = {"shares": float(res.shares_outstanding), "netDebt": float(-eb["B28"].value),
                      "adj": float(-eb["B29"].value), "leases": float(-eb["B30"].value)}
    return out, bridge


def _run_js(cfg: dict, expr: str):
    js = f"var CFG={json.dumps(cfg)};\n{_js_funcs()}\nconsole.log(JSON.stringify({expr}));"
    r = subprocess.run(["node", "-e", js], capture_output=True, text=True, timeout=60)
    assert r.returncode == 0, r.stderr
    return json.loads(r.stdout)


needs_node = pytest.mark.skipif(shutil.which("node") is None, reason="node not installed")


@needs_node
def test_m6_table_view_ev_equals_engine_ev_in_every_scenario():
    ev, bridge = _engine_pack()
    cfg = {"bridge": bridge, "engineEv": {n: v[0] for n, v in ev.items()}}
    # The UI carries the price to 3 decimals; the tie is within that rounding (AUD 1m).
    rows = [(n, round(v[1], 3)) for n, v in ev.items()]
    res = _run_js(cfg, "%s.map(function(r){var e=dcfBridgeEv(r[1]);return [r[0],e,dcfEngineTie(r[0],e)];})" % json.dumps(rows))
    assert len(res) == 6
    for name, table_ev, tie in res:
        assert tie["ok"], (name, tie)
        assert table_ev == pytest.approx(ev[name][0], abs=1.0)


@needs_node
def test_m6_tie_check_can_fail():
    """A tautological check cannot fail. Move the price 5% and this one must."""
    ev, bridge = _engine_pack()
    cfg = {"bridge": bridge, "engineEv": {n: v[0] for n, v in ev.items()}}
    n, (e, per) = "Muddle Through", ev["Muddle Through"]
    res = _run_js(cfg, "(function(){var x=dcfBridgeEv(%r*1.05);return dcfEngineTie(%r,x);})()" % (per, n))
    assert res["ok"] is False and abs(res["delta"]) > 100


def test_m6_scaffold_has_one_bridge_and_no_back_solved_tie():
    assert "ev=per*CFG.shares+CFG.netDebt" not in GEN_UI
    assert "reconstructed EV" not in GEN_UI
    assert "dcfEngineTie(sc.n, ev)" in GEN_UI
    assert "Less: equity-bridge adjustments (net)" in GEN_UI.split("function dcfTableHTML")[1]


def test_m6_cfg_bridge_and_engine_ev_are_the_engines_own():
    path = GEN / "cfgs_gen.json"
    if not path.exists():
        pytest.skip("cfgs_gen.json not generated (gitignored build output)")
    cfg = json.loads(path.read_text(encoding="utf-8"))["dnl"]
    ev, bridge = _engine_pack()
    for k, v in bridge.items():
        assert cfg["bridge"][k] == pytest.approx(v, rel=1e-9), k
    for n, (e, _per) in ev.items():
        assert cfg["engineEv"][n] == pytest.approx(e, rel=1e-9), n
    # The narrative view's headline EV is the same engine EV, to the rounding it prints.
    narr_ev = float(cfg["dcfRows"][0][1].replace(",", ""))
    assert narr_ev == pytest.approx(cfg["engineEv"]["Muddle Through"], abs=0.5)

"""D-86 reporting rules (a) and (b) in the generated workbooks.

(a) The goodwill-inclusive earned return sits beside the headline return as a
formula-driven disclosure. (b) CSL's starting-return and fade band sits beside
the headline value. Both are reporting only: the valuation rows do not read
them. The workbooks are recalculated in LibreOffice (independent of the Python
engine) and compared with the committed sets that pin the figures:

* ``design/methodology/terminal_return_sets.yaml`` (returns, all companies)
* ``design/methodology/csl_terminal_return_sets.yaml`` (CSL band)
"""

from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
import pytest
import yaml

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "ui_prototypes" / "_generator"))

import engine_workbook as ew  # noqa: E402

SOFFICE = shutil.which("soffice") or ("/usr/bin/soffice" if Path("/usr/bin/soffice").exists() else None)
RET_SET = yaml.safe_load((ROOT / "design/methodology/terminal_return_sets.yaml").read_text(encoding="utf-8"))
CSL_SET = yaml.safe_load((ROOT / "design/methodology/csl_terminal_return_sets.yaml").read_text(encoding="utf-8"))
ORDER = ["orderly_convergence", "muddle_through", "ai_productivity_lag", "fragmentation",
         "disorderly_climate_crystallisation", "stagflation_persists"]


def _rows(company):
    by = {r["scenario"]: r for r in RET_SET["terminal_return_sets"]["current"]["rows"] if r["company"] == company}
    return [by[s] for s in ORDER]


def _recalc(data: bytes, name: str):
    tmp = Path(tempfile.mkdtemp())
    src = tmp / f"{name}.xlsx"
    src.write_bytes(data)
    out = tmp / "out"
    subprocess.run([SOFFICE, "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=150)
    return openpyxl.load_workbook(out / f"{name}.xlsx", data_only=True)


@pytest.fixture(scope="module")
def built():
    return {"csl": ew.build_csl_workbook_bytes(), "dnl": ew.build_dnl_workbook_bytes()}


def test_disclosure_and_band_cells_are_formulas_and_goodwill_is_an_input(built):
    for company, data in built.items():
        import io
        wb = openpyxl.load_workbook(io.BytesIO(data))
        summary = wb["Scenarios"]
        for r in range(4, 10):
            for col in (5, 6) if company == "dnl" else (6, 7, 8, 9):
                assert str(summary.cell(r, col).value).startswith("="), (company, r, col)
        a = wb["Assumptions"]
        hits = [c for row in a.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("Goodwill") and "DISCLOSURE ONLY" in c.value]
        assert len(hits) == 1, company
        cell = a.cell(hits[0].row, 2)
        assert cell.fill.fill_type == "solid"                          # yellow input cell
        assert cell.font.color.rgb == ew.BLUEFONT.color.rgb            # blue text
        assert not str(cell.value).startswith("=")


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice not installed")
def test_csl_workbook_returns_and_band_tie_the_committed_sets(built):
    wb = _recalc(built["csl"], "csl")
    s = wb["Scenarios"]
    for i, row in enumerate(_rows("csl")):
        r = 4 + i
        assert round(s.cell(r, 6).value, 6) == pytest.approx(row["earned_return_ex_goodwill"], abs=2e-6)
        assert round(s.cell(r, 7).value, 6) == pytest.approx(row["earned_return_incl_goodwill"], abs=2e-6)
    band = wb["Terminal band"]
    for i, sc in enumerate(CSL_SET["scenarios"]):
        sens = {(b["fade_years"], b["starting_return"]): b["aud_per_share"] for b in sc["sensitivity"]}
        sens[(sc["fade_years"], "earned")] = sc["aud_per_share_model"]
        starts = sorted({k[1] for k in sens if k[1] != "earned"}, reverse=True)
        for j, fade in enumerate((sc["fade_years"], min(k[0] for k in sens))):
            r = 5 + 2 * i + j
            assert band.cell(r, 2).value == fade
            for c, key in enumerate(["earned"] + starts):
                assert band.cell(r, 3 + c).value == pytest.approx(sens[(fade, key)], abs=0.006), (sc["scenario"], fade, key)
        assert abs(band.cell(5 + 2 * i, 3 + 1 + len(starts)).value) < 1e-6   # model-fade check row
        # headline unchanged
        assert s.cell(4 + i, 3).value == pytest.approx(sc["aud_per_share_model"], abs=0.006)
        assert s.cell(4 + i, 8).value == pytest.approx(min(sens.values()), abs=0.006)
        assert s.cell(4 + i, 9).value == pytest.approx(max(sens.values()), abs=0.006)


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice not installed")
def test_dnl_workbook_returns_tie_the_committed_set(built):
    wb = _recalc(built["dnl"], "dnl")
    s = wb["Scenarios"]
    for i, row in enumerate(_rows("dnl")):
        r = 4 + i
        assert s.cell(r, 5).value == pytest.approx(row["earned_return_ex_goodwill"], abs=2e-6)
        assert s.cell(r, 6).value == pytest.approx(row["earned_return_incl_goodwill"], abs=2e-6)

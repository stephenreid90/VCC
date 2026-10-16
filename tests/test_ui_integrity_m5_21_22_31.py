"""Pins for UI tracker items M5, 21, 22 and 31 (16 Oct 2026).

M5/31  The DCF build-up panel is the engine's central case and does not move with the
       sliders; captions say so, and the override note says "approximation".
21     A "reset all" control clears this company's saved scenarios and world overrides.
22     DNL Equity bridge sheet carries no literal inputs; they sit on Assumptions.
CSL    The build-up rows are engine-sourced (no typed legacy figures) and describe the
       current WACC / excess-return-convergence method, not the old Ke / Gordon one.
"""

from __future__ import annotations

import contextlib
import io
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
GEN = ROOT / "ui_prototypes" / "_generator"
sys.path.insert(0, str(GEN))
sys.path.insert(0, str(ROOT / "src"))

import engine_workbook as ew  # noqa: E402

GEN_UI = (GEN / "gen_ui.py").read_text(encoding="utf-8")


@pytest.fixture(scope="module")
def bc():
    with contextlib.redirect_stdout(io.StringIO()):
        import build_cfgs
    return build_cfgs


def _texts(cfg):
    out = [cfg["dcf"], cfg["dcfIntro"], cfg["snap"]["dcf"], cfg["snap"]["discount"], cfg["_discount"]]
    for row in cfg["dcfRows"]:
        out.extend(str(x) for x in row)
    return out


def test_csl_rows_match_engine(bc):
    rows = {r[0].split(" (")[0]: r[1] for r in bc.csl["dcfRows"]}
    sp = bc._sp
    mt = "Muddle Through"
    assert rows["Sum PV explicit FCFF"] == "{:,.0f}".format(round(sp["pv_explicit"][mt]))
    assert rows["PV of terminal value"] == "{:,.0f}".format(round(sp["pv_terminal"][mt]))
    assert rows["Enterprise value"] == "{:,.0f}".format(round(sp["ev"][mt]))
    assert rows["Equity value"] == "{:,.0f}".format(round(sp["equity"][mt]))
    assert rows["Value per share USD"] == "%.2f" % sp["usd"][mt]
    assert rows["Value per share AUD"] == "%.2f" % sp["base"]
    assert rows["Less: net debt"] == "({:,.0f})".format(round(sp["net_debt"]))
    assert rows["÷ shares"] == "%.1f" % sp["shares"]


def test_csl_no_legacy_figures_or_method(bc):
    blob = " ".join(_texts(bc.csl))
    for legacy in ("74,029", "203.83", "134.52", "8.75% cost of equity", "Gordon-growth"):
        assert legacy not in blob, legacy
    rate = "%.2f%%" % (bc._sp["ke"] * 100)
    assert rate in bc.csl["dcfIntro"] and rate in bc.csl["snap"]["discount"]
    assert "excess-return convergence" in bc.csl["dcfIntro"]


def test_buildup_intros_say_they_do_not_move(bc):
    for cfg in (bc.dnl, bc.wbc, bc.csl):
        assert "does not move with the sliders" in cfg["dcfIntro"]
        assert "approximation" in cfg["dcfIntro"]


def test_no_straight_from_engine_caption(bc):
    blob = json.dumps({"dnl": bc.dnl, "wbc": bc.wbc, "csl": bc.csl}, ensure_ascii=False).lower()
    assert "straight from the production engine" not in blob
    assert "straight from the production engine" not in GEN_UI.lower()


def test_scaffold_has_reset_all_and_approximation():
    assert 'id="resetall"' in GEN_UI and "localStorage.removeItem(LS)" in GEN_UI
    assert "Reset done" in GEN_UI
    assert "reduced-form approximation" in GEN_UI
    for html in (ROOT / "ui_prototypes").glob("*_scenario_interface.html"):
        text = html.read_text(encoding="utf-8")
        assert 'id="resetall"' in text, html.name


def test_dnl_equity_bridge_has_no_literal_inputs():
    wb = load_workbook(io.BytesIO(ew.build_dnl_workbook_bytes()))
    ws = wb["Equity bridge"]
    bad = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                bad.append(f"{c.coordinate}={c.value}")
            if c.fill is not None and c.fill.fgColor is not None and c.fill.fgColor.rgb == ew.YELLOW.fgColor.rgb \
                    and c.value is not None:
                bad.append(f"{c.coordinate} input-styled")
            if isinstance(c.value, str) and c.value.startswith("=") and c.column > 1 and c.row > 2 \
                    and c.column in (2, 5, 6) and "Assumptions" not in c.value and "!" not in c.value \
                    and c.value.lstrip("=-").replace(".", "").isdigit():
                bad.append(f"{c.coordinate} constant formula")
    assert not bad, bad
    assump = wb["Assumptions"]
    labels = [str(c.value) for c in assump["A"] if c.value]
    assert any("probability" in x for x in labels) and any("— amount" in x for x in labels)


def test_market_price_sourced_from_data(bc):
    import yaml
    for short, cfg in (("csl", bc.csl), ("wbc", bc.wbc)):
        d = yaml.safe_load((ROOT / "data" / "companies" / f"{short}.yaml").read_text(encoding="utf-8"))
        px = d["company_position"]["share_statistics"]["market_reference_price_aud"]
        assert cfg["market"] == px
        assert cfg["mklab"] == "vs market (%.2f)" % px

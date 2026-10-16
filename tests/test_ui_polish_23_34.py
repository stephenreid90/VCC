"""Pins for UI tracker items 23, 24, 26, 27, 28, 29, 30, 33 and 34 (16 Oct 2026).

23  pvsub carries an engine-sourced gap sentence for every company.
24  Metric card 4 is the asymmetry for all three; terminal share sits in the discount-rate panel.
26  Bars show the delta vs Muddle Through beside the value.
27  Header download button; sheet count comes from the built workbook.
28  Provenance chip (commit, data date, build date).
29  Bars are keyboard operable and labelled; tabs carry aria-pressed and a non-colour cue.
30  Narrow widths: shrinking label column, value inside long bars, scrolling forces table.
33  First sentence plus the "why" line under the chart.
34  Dismissible three-step strip, remembered per company.
"""

from __future__ import annotations

import contextlib
import io
import json
import re
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
GEN = ROOT / "ui_prototypes" / "_generator"
sys.path.insert(0, str(GEN))
sys.path.insert(0, str(ROOT / "src"))

GEN_UI = (GEN / "gen_ui.py").read_text(encoding="utf-8")
KEYS = ("dnl", "wbc", "csl")


@pytest.fixture(scope="module")
def bc():
    with contextlib.redirect_stdout(io.StringIO()):
        import build_cfgs
    return build_cfgs


def _cfgs(bc):
    return {"dnl": bc.dnl, "wbc": bc.wbc, "csl": bc.csl}


def _html(key):
    return (ROOT / "ui_prototypes" / ("%s_scenario_interface.html" % key)).read_text(encoding="utf-8")


def test_pvsub_gap_sentence_all_companies(bc):
    for key, cfg in _cfgs(bc).items():
        assert cfg["pvsub"], key
        assert "%.2f" % cfg["market"] in cfg["pvsub"], key
        assert "%.2f" % cfg["cp"]["base"] in cfg["pvsub"], key
        assert ("above" in cfg["pvsub"]) or ("below" in cfg["pvsub"]), key
    assert "USD-functional" in bc.csl["pvsub"]  # CSL's existing text is kept


def test_metric4_is_asymmetry_for_all(bc):
    for key, cfg in _cfgs(bc).items():
        assert cfg["metric4"]["label"] == "Asymmetry (down/up)", key
        assert cfg["metric4"]["value"].endswith("×"), key
    assert bc.csl["metric4"]["value"] == "%.2f×" % bc._sp["asym"]


def test_terminal_share_in_discount_panel(bc):
    for key, cfg in _cfgs(bc).items():
        assert 'id="termshare"' in cfg["_discount"], key
        assert "11.4.2" in cfg["_discount"], key
    packs = {"dnl": bc._pk, "wbc": bc._wp, "csl": bc._sp}
    for key, pk in packs.items():
        assert 0.0 < pk["terminal_share"] < 1.0, key
        assert ("%.0f%% of %s" % (pk["terminal_share"] * 100, pk["terminal_denom"])) in _cfgs(bc)[key]["_discount"], key
    assert bc._wp["terminal_denom"] == "the equity claim"


def test_scaffold_hooks():
    for needle in ('id="hdl"', 'id="prov"', 'id="howto"', 'id="barnarr"', 'role="button" tabindex="0"',
                   "aria-pressed", "aria-label", "function updateBarNarr", "vcc_howto_", "ev.key==='Enter'"):
        assert needle in GEN_UI, needle
    assert "try{ localStorage.setItem(hk" in GEN_UI
    assert "textDecoration=on?'underline'" in GEN_UI


def test_delta_vs_mt_in_bar_renderer():
    assert "function fmtDelta" in GEN_UI
    assert "vs Muddle Through" in GEN_UI
    assert "(s.v/mtv-1)*100" in GEN_UI


def test_narrow_width_layout():
    assert not re.search(r"(?<!-)width:122px", GEN_UI)
    assert "flex:0 1 122px" in GEN_UI and "min-width:56px" in GEN_UI
    assert "left:calc('+(100-w)" not in GEN_UI and "right:calc('+(100-w)+'% + 6px)" in GEN_UI
    assert "<div style=\"overflow-x:auto;\"><table style=\"border-collapse:collapse;\">'+pr" in GEN_UI


def test_generated_pages_carry_the_features():
    for key in KEYS:
        h = _html(key)
        assert len(h) > 50000
        assert 'id="hdl"' in h and 'id="prov"' in h
        assert "Download audited workbook" in h
        assert "undefined" not in h.split("<script>")[0]


def test_generated_pages_match_current_cfgs(bc):
    for key, cfg in _cfgs(bc).items():
        h = _html(key)
        assert json.dumps(cfg["pvsub"])[1:-1] in h, key
        assert "termshare" in h, key


def test_xlsx_sheet_count_from_workbook(bc):
    import base64

    for key, cfg in _cfgs(bc).items():
        wb = load_workbook(io.BytesIO(base64.b64decode(cfg["xlsxB64"])), read_only=True)
        assert cfg["xlsxSheets"] == len(wb.sheetnames), key
        assert cfg["xlsxSheets"] > 1, key


def test_provenance(bc):
    for key, cfg in _cfgs(bc).items():
        p = cfg["prov"]
        assert p["commit"] and p["buildDate"], key
    assert bc.dnl["prov"]["dataDate"] is None
    assert bc.wbc["prov"]["dataDate"] and bc.csl["prov"]["dataDate"]


def test_csl_gap_claims_come_from_engine(bc):
    gap = "~%.0f%% above market" % (abs(bc._sp["discount_to_market"]) * 100.0)
    blob = " ".join(list(bc.csl["narr"].values()) + [bc.csl["footnote"], bc.csl["topnote"]] + list(bc.csl["snap"].values()))
    assert "93%" not in blob and "+52%" not in blob and "the whole 21% fall" not in blob
    assert gap in bc.csl["narr"]["Muddle Through"] and gap in bc.csl["footnote"]
    assert "just below market" not in bc.dnl["snap"]["world"]
    assert "~%.0f%% below market" % (abs(bc._pk["discount_to_market"]) * 100.0) in bc.dnl["snap"]["world"]


def test_csl_reads_as_wacc_wbc_keeps_cost_of_equity(bc):
    blob = json.dumps({k: v for k, v in bc.csl.items() if k != "xlsxB64"}, ensure_ascii=False)
    assert "cost-of-equity DCF" not in blob and "cost of equity)" not in blob
    assert bc.csl["ccynote"].endswith("WACC-based DCF") and bc.csl["discLabel"] == "WACC"
    assert [s["label"] for s in bc.csl["sliders"] if s["k"] == "re"] == ["Discount rate (WACC)"]
    assert "CFG.discLabel" in GEN_UI
    assert bc.wbc["beta"]["toDiscount"]["label"] == "cost of equity"
    assert "cost of equity" in [s["label"] for s in bc.wbc["sliders"] if s["k"] == "re"][0]
